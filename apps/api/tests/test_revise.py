import io
import re
import uuid

import pytest
from docx import Document
from docx.oxml.ns import qn

from legalai.db import session_scope
from legalai.drafting.revise import Item, plan
from legalai.ingestion.parsers import parse_document
from legalai.models import DocumentChunk, Redline
from tests.test_retrieval import drain, emb, mk_matter, upload  # noqa: F401
from tests.test_review import ReviewLLM

CLAUSES = [
    ("MASTER SERVICES AGREEMENT", "Heading 1", True),
    ("1.1 Definitions. Confidential Information means non-public information.", None, False),
    ("2.1 Term. This Agreement begins on the Effective Date and continues for two years.", None, False),
    ("(a) Either party may terminate for convenience on thirty days written notice.", None, False),
    ("(b) Either party may terminate immediately for uncured material breach.", None, False),
    ("3.1 Payment. Customer shall pay invoices within thirty days of receipt.", None, False),
]


def make_docx() -> bytes:
    d = Document()
    for text, style, bold in CLAUSES:
        p = d.add_paragraph(style=style) if style else d.add_paragraph()
        r = p.add_run(text)
        r.bold = bold or text.startswith("3.1")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def paras(data: bytes):
    return [(p.text, p.style.name, [r.bold for r in p.runs]) for p in Document(io.BytesIO(data)).paragraphs]


def text_variants(data: bytes):
    """(accept-all text, reject-all text) per paragraph, from the tracked-change XML (line breaks kept as newlines)."""
    d = Document(io.BytesIO(data))
    acc, rej = [], []
    for p in d.element.body.iter(qn("w:p")):
        a_parts, r_parts = [], []
        for el in p.iter():
            anc = {x.tag for x in el.iterancestors()}
            in_ins, in_del = qn("w:ins") in anc, qn("w:del") in anc
            if el.tag in (qn("w:t"), qn("w:delText")):
                txt = el.text or ""
            elif el.tag == qn("w:br"):
                txt = "\n"
            else:
                continue
            if not in_del:
                a_parts.append(txt)       # accepted view: everything except deletions
            if not in_ins:
                r_parts.append(txt)       # rejected view: everything except insertions
        acc.append("".join(a_parts))
        rej.append("".join(r_parts))
    return acc, rej


def start_flow(make_api, data=None, name="msa.docx", accept=True):
    a = make_api("admin@example.com")
    m = mk_matter(a, "standard")
    d = upload(a, m["id"], name, data or make_docx())
    drain()
    with session_scope() as db:
        chunk = db.query(DocumentChunk).filter(DocumentChunk.text.like("%terminate for convenience%")).first()
        cid, original = str(chunk.id), chunk.text
    r = a.post("/api/redlines", json={"chunk_id": cid, "instruction": "Extend the notice period", "perspective": "customer"})
    rid = r.json()["id"]
    drain(ReviewLLM())
    rl = a.get(f"/api/redlines/{rid}").json()
    assert rl["proposed"]
    if accept:
        a.patch(f"/api/redlines/{rid}", json={"status": "accepted"})
    return a, m, d, rid, original, rl["proposed"]


def test_in_place_docx_keeps_untouched_paragraphs_and_formatting(make_api, emb):
    a, m, d, rid, original, proposed = start_flow(make_api)
    orig_paras = paras(make_docx())
    r = a.post(f"/api/documents/{d['id']}/revise", json={"format": "docx"})
    assert r.status_code == 200 and r.headers["x-revise-applied"] == "1" and r.headers["content-disposition"].endswith('-revised.docx"')
    new = paras(r.content)
    assert new[0] == orig_paras[0] and new[0][1] == "Heading 1"                      # heading style untouched
    assert new[1] == orig_paras[1] and new[-1] == orig_paras[-1] and new[-1][2] == [True]   # bold run in a later paragraph untouched
    joined = " ".join(p[0] for p in new)
    assert "terminate for convenience on thirty days" not in joined or proposed.count("thirty") > 0
    assert proposed.split("\n\n")[0][:30] in joined and "uncured material breach" not in joined.replace(proposed, "")


def test_tracked_changes_accept_equals_revised_and_reject_equals_original(make_api, emb):
    a, m, d, rid, original, proposed = start_flow(make_api)
    r = a.post(f"/api/documents/{d['id']}/revise", json={"format": "tracked_docx"})
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('-tracked-changes.docx"')
    acc, rej = text_variants(r.content)
    src = [p[0] for p in paras(make_docx())]
    assert [t for t in rej if t.strip()] == src                                        # rejecting every change restores the original
    joined = " ".join(acc)
    assert proposed.replace("\n\n", " ")[:40].strip() in re.sub(r"\s+", " ", joined) and acc[0] == src[0]
    d2 = Document(io.BytesIO(r.content))
    assert len(d2.element.body.findall(".//" + qn("w:ins"))) >= 1 and len(d2.element.body.findall(".//" + qn("w:del"))) >= 1
    assert d2.element.body.find(".//" + qn("w:ins")).get(qn("w:author")) == "Legal AI"


def test_text_source_is_rebuilt_as_new_docx_and_txt_and_md(make_api, emb):
    contract = ("MASTER AGREEMENT\n\n2.1 Term. This Agreement continues for two years.\n\n(a) Either party may terminate for convenience on thirty days written notice.\n\n"
                "(b) Either party may terminate for uncured breach.\n\n3.1 Payment. Pay within thirty days.")
    a, m, d, rid, original, proposed = start_flow(make_api, contract.encode(), "msa.txt")
    pv = a.post(f"/api/documents/{d['id']}/revise/preview", json={}).json()
    assert pv["preserves_formatting"] is False and len(pv["applied"]) == 1 and pv["revised_chars"] > 0
    txt = a.post(f"/api/documents/{d['id']}/revise", json={"format": "txt"}).content.decode()
    assert txt.startswith("MASTER AGREEMENT") and txt.endswith("Pay within thirty days.") and original not in txt
    assert proposed.strip() in txt
    assert a.post(f"/api/documents/{d['id']}/revise", json={"format": "md"}).headers["content-disposition"].endswith('.md"')
    docx = a.post(f"/api/documents/{d['id']}/revise", json={"format": "docx"}).content
    assert any(p[0].endswith("Pay within thirty days.") for p in paras(docx)) and paras(docx)[0][0] == "MASTER AGREEMENT"
    tr = a.post(f"/api/documents/{d['id']}/revise", json={"format": "tracked_docx"}).content
    acc, rej = text_variants(tr)
    assert [t for t in rej if t.strip()] == [p for p in contract.split("\n\n")]


def test_nothing_applies_until_accepted_pending_can_be_included(make_api, emb):
    a, m, d, rid, original, proposed = start_flow(make_api, accept=False)
    assert a.post(f"/api/documents/{d['id']}/revise", json={"format": "txt"}).status_code == 409
    pv = a.post(f"/api/documents/{d['id']}/revise/preview", json={}).json()
    assert pv["applied"] == [] and pv["pending_not_applied"] == 1
    assert a.post(f"/api/documents/{d['id']}/revise", json={"format": "txt", "include_pending": True}).status_code == 200
    a.patch(f"/api/redlines/{rid}", json={"status": "rejected"})
    assert a.post(f"/api/documents/{d['id']}/revise", json={"format": "txt", "include_pending": True}).status_code == 409   # rejected never applies


def test_stale_redline_is_skipped_not_applied(make_api, emb):
    a, m, d, rid, original, proposed = start_flow(make_api)
    with session_scope() as db:
        db.get(Redline, uuid.UUID(rid)).original = original + " changed"
    assert a.post(f"/api/documents/{d['id']}/revise", json={"format": "txt"}).status_code == 409
    pv = a.post(f"/api/documents/{d['id']}/revise/preview", json={}).json()
    assert pv["applied"] == [] and "changed since" in pv["skipped"][0]["reason"]


def test_save_as_new_version_is_reingested_and_searchable(make_api, emb):
    a, m, d, rid, original, proposed = start_flow(make_api)
    r = a.post(f"/api/documents/{d['id']}/revise", json={"format": "docx", "save_as_version": True})
    assert r.status_code == 200 and r.json()["saved"] and r.json()["version_number"] == 2
    drain()
    vs = a.get(f"/api/documents/{d['id']}/versions").json()
    assert [v["version_number"] for v in vs] == [1, 2] and vs[1]["parse_status"] == "parsed" and vs[1]["filename"].endswith("-revised.docx")
    doc = a.get(f"/api/documents/{d['id']}").json()
    assert doc["current_version"]["version_number"] == 2
    cmp = a.post("/api/comparisons", json={"version_a": vs[0]["id"], "version_b": vs[1]["id"]}).json()
    assert cmp["counts"]["modified"] + cmp["counts"]["added"] + cmp["counts"]["removed"] >= 1
    acts = {e["action"] for e in a.get(f"/api/matters/{m['id']}/audit").json()}
    assert "document.revised" in acts


def test_manual_recommendations_listed_from_latest_review(make_api, emb):
    a, m, d, rid, original, proposed = start_flow(make_api)
    a.post("/api/contracts/review", json={"document_id": d["id"]})
    drain(ReviewLLM())
    pv = a.post(f"/api/documents/{d['id']}/revise/preview", json={}).json()
    assert any(i["basis"] == "missing_term" for i in pv["needs_manual_drafting"])


def test_revise_permissions(make_api, emb):
    a, m, d, rid, original, proposed = start_flow(make_api)
    bob = make_api("bob@example.com")
    assert bob.post(f"/api/documents/{d['id']}/revise", json={"format": "txt"}).status_code == 404
    assert bob.post(f"/api/documents/{d['id']}/revise/preview", json={}).status_code == 404


def test_plan_verifies_offsets_and_overlaps():
    text = "AAA one BBB two CCC three"
    ok = plan(text, [Item("1", "x", "one", "ONE", 4, 7), Item("2", "y", "BBB", "bbb", 8, 11)])
    assert ok.revised_text == "AAA ONE bbb two CCC three" and len(ok.applied) == 2
    bad = plan(text, [Item("1", "x", "WRONG", "z", 4, 7), Item("2", "y", "one BBB", "q", 4, 11), Item("3", "z", "BBB", "w", 8, 11)])
    assert [s["redline_id"] for s in bad.skipped] == ["1"] and bad.revised_text == "AAA q two CCC three" or len(bad.applied) == 1
    assert plan(text, [Item("1", "x", "one", "  ", 4, 7)]).skipped[0]["reason"] == "No proposed text"
