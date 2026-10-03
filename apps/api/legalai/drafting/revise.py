"""Recreate an uploaded document with its recommendations applied.

Inputs are redlines (always: original = the stored clause text, proposed = replacement). Output formats:
  txt / md        revised plain text
  docx            clean revised Word document
  tracked_docx    Word document with real tracked changes (w:ins / w:del) that you can accept or reject in Word
For originals that are .docx the ORIGINAL FILE is edited in place, so every untouched paragraph keeps its styles and
formatting. Other formats are rebuilt as a new Word document from the extracted text (original layout is not preserved).
Every change is verified against the stored text first: a redline whose original no longer matches is skipped and
reported, never applied blindly."""
import copy
import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone

from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..ingestion.parsers import _para_text
from ..review.diff import word_diff

AUTHOR = "Legal AI"


@dataclass
class Item:
    redline_id: str
    label: str
    original: str
    proposed: str
    start: int
    end: int


@dataclass
class Result:
    applied: list[Item] = field(default_factory=list)
    skipped: list[dict] = field(default_factory=list)
    revised_text: str = ""


def plan(full_text: str, items: list[Item]) -> Result:
    """Verify each item against the stored text, drop overlaps, produce the revised plain text."""
    res = Result()
    ok: list[Item] = []
    for it in sorted(items, key=lambda i: i.start):
        if full_text[it.start:it.end] != it.original:
            res.skipped.append({"redline_id": it.redline_id, "label": it.label, "reason": "The document text changed since this redline was made"})
        elif not it.proposed.strip():
            res.skipped.append({"redline_id": it.redline_id, "label": it.label, "reason": "No proposed text"})
        elif ok and it.start < ok[-1].end:
            res.skipped.append({"redline_id": it.redline_id, "label": it.label, "reason": "Overlaps another change"})
        else:
            ok.append(it)
    out, pos = [], 0
    for it in ok:
        out += [full_text[pos:it.start], it.proposed.strip()]
        pos = it.end
    out.append(full_text[pos:])
    res.applied, res.revised_text = ok, "".join(out)
    return res


# ---- shared docx helpers ----------------------------------------------------------------------------------------
_ID = [1000]


def _stamp(el, author=AUTHOR):
    _ID[0] += 1
    el.set(qn("w:id"), str(_ID[0]))
    el.set(qn("w:author"), author)
    el.set(qn("w:date"), datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))


def _runs(kind: str, text: str, rpr):
    """Word run elements for `text` (newlines become line breaks). kind: equal | insert | delete."""
    out = []
    for i, part in enumerate(text.split("\n")):
        r = OxmlElement("w:r")
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        if i:
            r.append(OxmlElement("w:br"))
        if part:
            t = OxmlElement("w:delText" if kind == "delete" else "w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = part
            r.append(t)
        out.append(r)
    if kind == "equal":
        return out
    wrap = OxmlElement("w:ins" if kind == "insert" else "w:del")
    _stamp(wrap)
    for r in out:
        wrap.append(r)
    return [wrap]


def _fill(p_el, segments, rpr):
    for child in list(p_el):
        if child.tag != qn("w:pPr"):
            p_el.remove(child)
    for kind, text in segments:
        for el in _runs(kind, text, rpr):
            p_el.append(el)


def _first_rpr(p_el):
    r = p_el.find(qn("w:r"))
    return r.find(qn("w:rPr")) if r is not None else None


def _norm(t: str) -> str:
    return re.sub(r"[ \t ]+", " ", t).strip()


def _segments_for_block(block_text: str, bstart: int, pieces: list[tuple[int, int, str, str]], tracked: bool):
    """Segments for one source paragraph. pieces: (start, end, original, proposed) overlapping this block, in order."""
    segs, pos = [], 0
    for s, e, orig, prop in pieces:
        ls, le = max(s - bstart, 0), min(e - bstart, len(block_text))
        if ls > pos:
            segs.append(("equal", block_text[pos:ls]))
        first = s >= bstart  # the proposed text is inserted once, in the paragraph where the change starts
        if tracked:
            if le > ls:
                segs.append(("delete", block_text[ls:le]))
            if first:
                segs.append(("insert", prop.replace("\n\n", "\n")))
        elif first:
            segs.append(("insert", prop))
        pos = max(pos, le)
    if pos < len(block_text):
        segs.append(("equal", block_text[pos:]))
    return segs


# ---- edit the original DOCX in place -------------------------------------------------------------------------------
def revise_original_docx(orig: bytes, applied: list[Item], tracked: bool) -> bytes:
    doc = DocxDocument(io.BytesIO(orig))
    blocks = []  # (p_element, text, start, end) aligned with the parser's offsets
    pos = 0
    for p in doc.element.body.iter(qn("w:p")):
        t = _norm(_para_text(p))
        if t:
            blocks.append((p, t, pos, pos + len(t)))
            pos += len(t) + 2
    for p, text, bs, be in blocks:
        pieces = [(it.start, it.end, it.original, it.proposed) for it in applied if it.start < be and it.end > bs]
        if not pieces:
            continue
        # a change that starts in an earlier paragraph also consumes the start of this one
        segs = _segments_for_block(text, bs, pieces, tracked)
        rpr = _first_rpr(p)
        if not tracked:
            paras = [[]]
            for kind, t in segs:
                parts = t.split("\n\n")
                paras[-1].append((kind, parts[0]))
                for extra in parts[1:]:
                    paras.append([(kind, extra)])
            first = [s_ for s_ in paras[0] if s_[1]]
            if not first and len(paras) == 1:  # paragraph fully consumed by a change that began earlier
                p.getparent().remove(p)
                continue
            _fill(p, first or [("equal", "")], rpr)
            anchor = p
            for extra in paras[1:]:
                new_p = copy.deepcopy(p)
                _fill(new_p, [s for s in extra if s[1]] or [("equal", "")], rpr)
                anchor.addnext(new_p)
                anchor = new_p
        else:
            _fill(p, segs, rpr)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ---- rebuild as a new DOCX from text -------------------------------------------------------------------------------
HEADING = re.compile(r"^(ARTICLE\s+[IVXLC\d]+\b.*|[A-Z][A-Z0-9 ,.&'\-]{5,}|\d+(\.\d+)*\.?\s+[A-Z][^.]{0,60})$")


def _new_doc() -> DocxDocument:
    d = DocxDocument()
    return d


def build_clean_docx(text: str, title: str = "") -> bytes:
    d = _new_doc()
    if title:
        d.core_properties.title = title
    for para in text.split("\n\n"):
        p = d.add_paragraph()
        run = p.add_run(para)
        if len(para) < 90 and HEADING.match(para.strip()):
            run.bold = True
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def build_tracked_docx(full_text: str, applied: list[Item], title: str = "") -> bytes:
    d = _new_doc()
    if title:
        d.core_properties.title = title
    seq: list[tuple[str, str]] = []
    pos = 0
    for it in applied:
        if it.start > pos:
            seq.append(("equal", full_text[pos:it.start]))
        for op in word_diff(it.original, it.proposed.strip()):
            # inserted text uses line breaks, so paragraph structure always follows the ORIGINAL text and
            # rejecting every change restores the original document exactly
            seq.append((op["op"], op["text"].replace("\n\n", "\n") if op["op"] == "insert" else op["text"]))
        pos = it.end
    seq.append(("equal", full_text[pos:]))
    paras: list[list[tuple[str, str]]] = [[]]
    for kind, text in seq:
        parts = text.split("\n\n")
        if parts[0]:
            paras[-1].append((kind, parts[0]))
        for extra in parts[1:]:
            paras.append([(kind, extra)] if extra else [])
    for segs in paras:
        p = d.add_paragraph()
        for kind, text in segs:
            for el in _runs(kind, text, None):
                p._p.append(el)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def render(fmt: str, *, full_text: str, res: Result, original: bytes | None, original_name: str, title: str) -> tuple[bytes, str, str]:
    """-> (file bytes, filename, media type)"""
    stem = original_name.rsplit(".", 1)[0]
    docx_mt = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    is_docx = original_name.lower().endswith(".docx") and original is not None
    if fmt in ("txt", "md"):
        return res.revised_text.encode(), f"{stem}-revised.{fmt}", "text/plain" if fmt == "txt" else "text/markdown"
    if fmt == "docx":
        data = revise_original_docx(original, res.applied, False) if is_docx else build_clean_docx(res.revised_text, title)
        return data, f"{stem}-revised.docx", docx_mt
    if fmt == "tracked_docx":
        data = revise_original_docx(original, res.applied, True) if is_docx else build_tracked_docx(full_text, res.applied, title)
        return data, f"{stem}-tracked-changes.docx", docx_mt
    raise ValueError(f"Unknown format {fmt}")
